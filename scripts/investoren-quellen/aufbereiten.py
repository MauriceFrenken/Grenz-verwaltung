"""
Externe Investorenlisten (7Plus, Leadliste, Suchkunden) normalisieren und mit Salesforce abgleichen.

Aufruf:
  python3 -I scripts/investoren-quellen/aufbereiten.py QUELLORDNER SF_JSON NRW_ORTE_JSON AUSGABE_JSON

QUELLORDNER   Ordner mit Investorenliste_7Plus.xlsx, Investoren_7Plus_Teil_1.xlsx,
              Leadliste_Investoren_Grenz.xlsx, Suchkunden.xlsx
SF_JSON       Liste {id, email, acc, firma} aller Salesforce-Kontakte und -Leads mit E-Mail
NRW_ORTE_JSON Liste der NRW-Orte (Kreise, Städte, Gemeinden) aus den Ankaufsprofilen
AUSGABE_JSON  Zeilen für die Tabelle investor_quellen

Die Ausgabe enthält personenbezogene Kontaktdaten und gehört nicht ins Repository.
"""

import collections
import json
import re
import sys
import unicodedata

import openpyxl

ORDNER, SF_DATEI, NRW_DATEI, AUSGABE = sys.argv[1:5]

FREEMAIL = {
    "gmail.com", "gmx.de", "web.de", "t-online.de", "googlemail.com", "hotmail.com", "yahoo.de", "yahoo.com",
    "outlook.com", "outlook.de", "icloud.com", "gmx.net", "aol.com", "freenet.de", "mail.de", "me.com", "live.de",
    "arcor.de", "posteo.de", "hotmail.de", "online.de", "gmx.at", "bluewin.ch",
}
RECHTSFORM = re.compile(
    r"\b(gmbh|mbh|ag|kg|kgaa|se|ug|ohg|gbr|eg|ev|co|ltd|llc|inc|plc|sarl|sa|bv|nv|holding|gruppe|group|"
    r"invest(?:ment)?s?|immobilien|real estate|asset management|capital|partners?|haftungsbeschrankt)\b"
)


def firmen_schluessel(s):
    s = unicodedata.normalize("NFKD", str(s or "").lower())
    s = "".join(c for c in s if not unicodedata.combining(c)).replace("&", " ")
    s = RECHTSFORM.sub(" ", s)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def mails(s):
    return sorted(set(re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", str(s or "").lower())))


def domain(m):
    d = m.split("@")[1]
    return None if d in FREEMAIL else d


def text(v):
    if v is None:
        return None
    t = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()
    return t or None


def mio(v):
    m = re.search(r"([\d.,]+)\s*Mio", str(v or ""))
    return round(float(m.group(1).replace(".", "").replace(",", ".")) * 1_000_000) if m else None


def anzahl(v):
    m = re.search(r"(\d+)", str(v or ""))
    return int(m.group(1)) if m else None


def liste(v):
    return [t for t in (text(x) for x in str(v or "").split(";")) if t]


TYP_REGELN = [
    (r"multi family office|family office", "Family Office"),
    (r"fonds|kvg|asset manager", "Fonds/KVG"),
    (r"versicherung|institutionell", "Institutionell/Versicherung"),
    (r"entwickler", "Projektentwickler"),
    (r"privat", "Privat"),
]


def typ_aus(v):
    w = str(v or "").lower()
    for re_, typ in TYP_REGELN:
        if re.search(re_, w):
            return typ
    return None


NUTZUNG = [
    (r"wohn- / gesch|wohn-/ gesch|geschäftsh", "Wohn-/ Geschäftshaus"),
    (r"mehrfamilien|wohnungs|wohnen|stadtquartier", "Wohnen"),
    (r"büro|office", "Büro"),
    (r"fachm|supermarkt|discounter|warenh|nahversorg|handel|einkauf", "Einzelhandel"),
    (r"logistik|lager|produktion|gewerbeh|industrial", "Logistik/Gewerbe"),
    (r"hotel", "Hotel"),
    (r"pflege|healthcare|ärzte|klinik", "Pflege"),
    (r"student", "Studentenwohnen"),
    (r"micro|serviced", "Mirco-Apartments"),
    (r"grundst", "Grundstück"),
]


def nutzungen(*werte):
    gefunden = []
    for w in werte:
        for teil in liste(w) or ([text(w)] if text(w) else []):
            for re_, n in NUTZUNG:
                if re.search(re_, teil.lower()) and n not in gefunden:
                    gefunden.append(n)
    return gefunden


def zeilen(datei, blatt=None, kopf=0):
    wb = openpyxl.load_workbook(f"{ORDNER}/{datei}", read_only=True, data_only=True)
    ws = wb[blatt] if blatt else wb.worksheets[0]
    alle = list(ws.iter_rows(values_only=True))[kopf:]
    namen = [str(h).strip() if h is not None else f"_{i}" for i, h in enumerate(alle[0])]
    # Doppelte Spaltennamen (Teil 1 hat Objektvolumen zweimal): erste nicht leere gewinnt.
    ergebnis = []
    for r in alle[1:]:
        d = {}
        for k, v in zip(namen, r):
            if v not in (None, "") and k not in d:
                d[k] = v
        ergebnis.append(d)
    return ergebnis


# Salesforce-Abgleich
sf = json.load(open(SF_DATEI))
sf_mail = {e["email"].lower(): e for e in sf if e.get("email")}
sf_domain = collections.defaultdict(list)
sf_name = {}
for e in sf:
    if e.get("email") and domain(e["email"].lower()):
        sf_domain[domain(e["email"].lower())].append(e)
    n = firmen_schluessel(e.get("firma"))
    if len(n) >= 4 and n != "privat":
        sf_name.setdefault(n, e)


def sf_treffer(mailliste, firma):
    for m in mailliste:
        if m in sf_mail:
            return "email", sf_mail[m]
    n = firmen_schluessel(firma)
    if len(n) >= 4 and n in sf_name:
        return "name", sf_name[n]
    for m in mailliste:
        d = domain(m)
        if d and d in sf_domain:
            return "domain", sf_domain[d][0]
    return None, None


NRW = {o.lower() for o in json.load(open(NRW_DATEI))} | {"nordrhein-westfalen", "nrw", "rheinland", "ruhrgebiet"}
NRW_PLZ = re.compile(r"\b(4[0-7]\d{3}|5[0-3]\d{3}|57[0-5]\d{2}|58\d{3}|59[0-5]\d{2}|32\d{3}|33[0-4]\d{2})\b")


def nrw_bezug(regionen, kriterium, sitz):
    # Ortsnamen gibt es mehrfach (Berg, Altdorf): erst mehrere NRW-Treffer mit nennenswertem Anteil zählen.
    treffer = sum(1 for r in regionen if r.lower() in NRW)
    if treffer >= 3 and treffer >= 0.15 * len(regionen):
        return "region"
    if re.search(r"nordrhein|deutschland|bundesweit", str(kriterium or ""), re.I):
        return "region"
    if NRW_PLZ.search(str(sitz or "")):
        return "sitz"
    return None


def passung(z):
    """0 bis 100: wie gut passt der Investor zu Grenz (NRW, Wohnen/WGH, 1 bis 20 Mio. €, erreichbar)."""
    p = 0
    p += {"region": 40, "sitz": 25}.get(z["nrw_bezug"], 0)
    if {"Wohnen", "Wohn-/ Geschäftshaus"} & set(z["nutzungen"]):
        p += 25
    vmin, vmax = z["volumen_min"] or 0, z["volumen_max"] or 10**12
    if vmin <= 20_000_000 and vmax >= 1_000_000:
        p += 20
    if z["emails"]:
        p += 10
    if z["typ"]:
        p += 5
    return min(p, 100)


def basis(quelle, schluessel, firma, person, mailliste):
    art, t = sf_treffer(mailliste, firma)
    return {
        "quelle": quelle,
        "schluessel": f"{quelle}:{schluessel}",
        "firma": text(firma),
        "person": text(person),
        "emails": mailliste,
        "sf_treffer_art": art,
        "sf_kontakt_sf_id": t["id"] if t else None,
        "sf_account_sf_id": (t.get("acc") if t else None),
        "typ": None, "nutzungen": [], "objektklassen": [], "zustand": [], "struktur": [],
        "volumen_min": None, "volumen_max": None, "einheiten_min": None, "einheiten_max": None,
        "regionen": [], "nrw_bezug": None, "objekte_interesse": [], "nicht_anlegen": False,
        "telefon": None, "adresse": None, "ort": None, "webseite": None, "quelle_link": None,
        "roh": {},
    }


ausgabe = {}

# 7Plus: Teil 1 (Detailprofile) und Listings (Kurzprofile), je Firma + Assetklasse ein Datensatz.
for t in zeilen("Investoren_7Plus_Teil_1.xlsx", kopf=1):
    firma = t.get("listing_company_name") or t.get("title")
    if not firma:
        continue
    klasse = text(t.get("klasse")) or text(t.get("listing_type"))
    z = basis("7plus", f"{firmen_schluessel(firma)}:{klasse}", firma, t.get("ansprechpartner"), mails(t.get("e-mail adresse")))
    regionen = liste(re.sub(r"^.*?(?=[A-ZÄÖÜ][a-zäöüß]+;)", "", str(t.get("staedte") or ""), count=1))
    z.update(
        typ=typ_aus(t.get("listing_investor_type") or t.get("type")),
        nutzungen=nutzungen(t.get("Nutzungsart"), klasse),
        zustand=liste(t.get("Objektzustand")),
        struktur=liste(t.get("struktur")),
        objektklassen=liste(t.get("objektklassen")),
        volumen_min=mio(t.get("Objektvolumen Min.")) or mio(t.get("listing_min_investment")),
        volumen_max=mio(t.get("Objektvolumen Max.")) or mio(t.get("listing_max_investment")),
        einheiten_min=anzahl(t.get("Anzahl Einheiten Min.")),
        einheiten_max=anzahl(t.get("Anzahl Einheiten Max.")),
        regionen=regionen[:600],
        telefon=text(t.get("telefonnummer")),
        adresse=text(t.get("sitz")),
        webseite=text(t.get("webseite")),
        quelle_link=text(t.get("listing_href")),
        roh={"klasse": klasse, "kriterium": text(t.get("kriterium")), "gewerbeschein": text(t.get("gewerbeschein"))},
    )
    z["nrw_bezug"] = nrw_bezug(regionen, t.get("kriterium"), t.get("sitz"))
    ausgabe.setdefault(z["schluessel"], z)

for r in zeilen("Investorenliste_7Plus.xlsx", "Listings"):
    firma = r.get("Company Name")
    if not firma:
        continue
    klasse = text(r.get("Listing Type"))
    z = basis("7plus", f"{firmen_schluessel(firma)}:{klasse}", firma, r.get("Ansprechpartner"), mails(r.get("E-Mail-Adressse")))
    if z["schluessel"] in ausgabe:
        continue
    z.update(
        typ=typ_aus(r.get("Investor Type")),
        nutzungen=nutzungen(klasse),
        volumen_min=mio(r.get("Minimum Investment")),
        volumen_max=mio(r.get("Maximum Investment")),
        telefon=text(r.get("Telefonnummer Festnetz")) or text(r.get("Telefonnummer Mobil")),
        adresse=text(r.get("Adresse")),
        webseite=text(r.get("Webseite")),
        quelle_link=text(r.get("Link")),
        roh={"klasse": klasse, "ankaufsprofil": text(r.get("Link Ankaufsprofil"))},
    )
    z["nrw_bezug"] = nrw_bezug([], None, r.get("Adresse"))
    ausgabe[z["schluessel"]] = z

# Leadliste Investoren Grenz
for r in zeilen("Leadliste_Investoren_Grenz.xlsx", kopf=1):
    firma = r.get("Company Name")
    if not firma:
        continue
    z = basis("leadliste", firmen_schluessel(firma), firma, r.get("Ansprechpartner"), mails(r.get("E-Mail-Adressse")))
    z.update(
        typ=typ_aus(r.get("Investor Type")),
        telefon=text(r.get("Telefonnummer Festnetz")) or text(r.get("Telefonnummer Mobil")),
        adresse=text(r.get("Adresse")),
        quelle_link=text(r.get("Link Ankaufsprofil")),
    )
    z["nrw_bezug"] = nrw_bezug([], None, r.get("Adresse"))
    ausgabe[z["schluessel"]] = z

# Suchkunden: eine Zeile je Person, Objekte mit Interesse gesammelt.
kopf = zeilen("Suchkunden.xlsx", "Alles")
objekt_spalten = [k for k in (kopf[0].keys() if kopf else []) if re.match(r"^\d+ ", k)]
wb = openpyxl.load_workbook(f"{ORDNER}/Suchkunden.xlsx", read_only=True, data_only=True)
roh = list(wb["Alles"].iter_rows(values_only=True))
namen = [str(h).strip() if h else "" for h in roh[0]]
objekt_spalten = [i for i, n in enumerate(namen) if re.match(r"^\d+ ", n)]
for r in roh[1:]:
    d = dict(zip(namen, r))
    vor, nach = text(d.get("Vorname")), text(d.get("Nachname"))
    if not (vor or nach):
        continue
    m = mails(d.get("E-Mail (P)"))
    key = m[0] if m else firmen_schluessel(f"{vor} {nach}")
    z = ausgabe.get(f"suchkunden:{key}")
    if not z:
        z = basis("suchkunden", key, d.get("Firmenname (F)"), " ".join(x for x in (vor, nach) if x), m)
        z.update(
            telefon=text(d.get("Telefon (Mobil) (P)")) or text(d.get("Telefon (geschäftl.) (P)")) or text(d.get("Telefon (F)")),
            adresse=" ".join(str(x) for x in (d.get("Straße (P)"), d.get("PLZ (P)"), d.get("Ort (P)")) if x),
            ort=text(d.get("Ort (P)")),
        )
        z["nrw_bezug"] = nrw_bezug([], None, z["adresse"])
        z["typ"] = "Privat" if not z["firma"] else None
        ausgabe[z["schluessel"]] = z
    for i in objekt_spalten:
        if r[i] not in (None, "") and namen[i] not in z["objekte_interesse"]:
            z["objekte_interesse"].append(namen[i])
    if str(d.get("In Salesforce") or "").strip().lower() == "nicht anlegen":
        z["nicht_anlegen"] = True
    if str(d.get("Adresse aktiv") or "").strip().lower() == "nein":
        z["roh"]["adresse_inaktiv"] = True
    if str(d.get("Status") or "").strip().lower().startswith("inakt"):
        z["roh"]["status"] = "inaktiv"

# Suchkunden haben Interesse an Objekten im Rheinland gezeigt: das zählt als NRW-Bezug.
for z in ausgabe.values():
    if z["quelle"] == "suchkunden" and z["objekte_interesse"]:
        z["nrw_bezug"] = z["nrw_bezug"] or "region"
    z["passung"] = passung(z)

json.dump(list(ausgabe.values()), open(AUSGABE, "w"), ensure_ascii=False, default=str)
zaehler = collections.Counter((z["quelle"], bool(z["sf_treffer_art"])) for z in ausgabe.values())
print(len(ausgabe), dict(zaehler))
print("Passung >= 70:", collections.Counter(z["quelle"] for z in ausgabe.values() if z["passung"] >= 70 and not z["sf_treffer_art"]))
